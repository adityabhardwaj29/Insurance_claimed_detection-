"""
build_viva_docx.py - Master Orchestrator Script for PROJECT_VIVA_PREPARATION_GUIDE.docx
Assembles all 41 sections across Parts 1, 2, 3, and 4 into a beautifully
formatted, professional Word (.docx) document.
"""

import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

# Ensure scripts dir is on sys.path
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
if SCRIPT_DIR not in sys.path:
    sys.path.insert(0, SCRIPT_DIR)

from viva_helpers import (
    setup_document, add_title, add_heading_1, add_heading_2,
    add_body, add_callout, add_table_data, COLOR_NAVY, COLOR_MUTED
)
from viva_part1 import build_part1
from viva_part2 import build_part2
from viva_part3 import build_part3
from viva_part4 import build_part4

def build_cover_page(doc: Document):
    """Creates a prestigious academic cover page for the Viva Guide."""
    p_pre = doc.add_paragraph()
    p_pre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_pre = p_pre.add_run("ACADEMIC VIVA DEFENSE PREPARATION HANDBOOK\nBACHELOR OF DATA SCIENCE (SEMESTER V, 2026–2027)")
    run_pre.font.name = "Times New Roman"
    run_pre.font.size = Pt(11)
    run_pre.font.bold = True
    run_pre.font.color.rgb = COLOR_MUTED
    p_pre.paragraph_format.space_before = Pt(36)
    p_pre.paragraph_format.space_after = Pt(28)

    # Main Project Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("GRAPH-ENHANCED INSURANCE CLAIM ANOMALY AND DUPLICATE NETWORK DETECTION")
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_NAVY
    p_title.paragraph_format.space_after = Pt(14)

    # Subtitle / Platform Name
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Platform: FraudShield AI — Multi-Signal Intelligence Decision Support System\nComplete Oral Defense Manual, Technical Explanations, Empirical Verification & 180+ Viva Questions")
    run_sub.font.name = "Times New Roman"
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    p_sub.paragraph_format.space_after = Pt(40)

    # Student & Evaluation Metadata Box / Table
    meta_table_data = [
        ["Attribute", "Academic Specification"],
        ["Candidate Name", "Aditya Bhardwaj"],
        ["Programme", "Bachelor of Data Science (B.Sc. Data Science)"],
        ["Academic Year / Sem", "Third Year | Semester V | Academic Year: 2026 – 2027"],
        ["Roll Number / Div", "TDDS003A | Division A"],
        ["Project Guide", "Ms. Sweta Suman (Assistant Professor)"],
        ["Department / School", "Department of Data Science"],
        ["Primary Evaluation Focus", "Tabular ML, Unsupervised Anomaly, Graph Theory, Duplicate NLP, XAI"],
        ["Empirical Benchmark", "320 Claims (52 Fraud / 16.25%), 1,020 Nodes, 2,615 Edges"],
        ["Language of Defense", "Simple Hinglish (English Technical Terminology + Hindi Spoken Flow)"]
    ]
    add_table_data(doc, meta_table_data, [2.2, 4.8])

    # Decorative bottom note
    p_bot = doc.add_paragraph()
    p_bot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_bot = p_bot.add_run("\nCONFIDENTIAL ACADEMIC VIVA PREPARATION RESOURCE\nSTRICTLY ZERO FABRICATION — 100% REPOSITORY VERIFIED METRICS")
    run_bot.font.name = "Times New Roman"
    run_bot.font.size = Pt(10)
    run_bot.font.bold = True
    run_bot.font.color.rgb = RGBColor(0x05, 0x96, 0x69) # Emerald Green
    p_bot.paragraph_format.space_before = Pt(36)
    p_bot.paragraph_format.space_after = Pt(24)

    doc.add_page_break()

def build_table_of_contents_summary(doc: Document):
    """Creates a comprehensive section navigation index."""
    add_heading_1(doc, "TABLE OF CONTENTS & HANDBOOK ROADMAP")
    
    add_body(doc, "Ye preparation guide 41 comprehensive sections me structured hai jo basics se lekar high-level architecture tak cover karta hai:")

    toc_data = [
        ["Section Range", "Theme & Modules Covered", "Primary Viva Focus"],
        ["Section 1 – 5", "Project at a Glance, 30s-5m Speeches, Problem Statement, Existing vs Proposed System.", "High-level understanding, elevator pitch, motivation."],
        ["Section 6 – 10", "12-Step Intake-to-Triage Workflow, ML Algorithms, Supervised vs Unsupervised, XGBoost Deep Dive, Anomaly Detection.", "Mathematical algorithms, loss functions, parameter tuning."],
        ["Section 11 – 15", "Duplicate Detection (SequenceMatcher, Jaccard, TF-IDF), Graph Analysis (NetworkX), Graph vs RDBMS, Composite Risk Formula, Explainable AI (SHAP).", "Core differentiators, graph topology, multi-signal fusion."],
        ["Section 16 – 20", "Dataset Verification (320 claims), Temporal Train/Test Split (224/96), PR-AUC Evaluation, Confusion Matrix (N=96), Class Imbalance Handling.", "Empirical defense, zero fabrication, anti-leakage design."],
        ["Section 21 – 25", "Tech Stack Breakdown, React 18 Frontend, FastAPI Backend, PostgreSQL Supabase Database, Authentication & Security.", "Full-stack software engineering, REST architecture."],
        ["Section 26 – 30", "14 REST APIs Specification, Cloud Deployment (Vercel/Render), Testing Suite, System Limitations, Future Scope.", "Production readiness, scalability, honest self-critique."],
        ["Section 31", "105 Most Important Viva Questions in 5 Levels (Basic, Project, ML, Advanced, Trick Questions).", "Question-answer preparation with 4-part structure."],
        ["Section 32 – 35", "50 Rapid Fire Questions, 14 'Why' Questions, 10 'What If' Scenarios, 17 Examiner Challenge Questions.", "Quick response drills, unexpected situation defense."],
        ["Section 36 – 41", "Project Journey, Personal Contribution, Live Demo Script, 1-Page Revision Sheet, 10 Must-Memorize Answers, Quality Checklist.", "Live presentation mastery, confidence building."]
    ]
    add_table_data(doc, toc_data, [1.4, 3.4, 2.2])

    add_callout(doc, "HOW TO STUDY THIS GUIDE",
        "Aditya, is document ko padhne ka best formula:\n"
        "1. Step 1: Section 1 ke charo speech formats (30s, 1m, 3m, 5m) ko loud voice me bolkar memorize karein.\n"
        "2. Step 2: Section 16 se 20 ke factual metrics (320 claims, 16.25% fraud, 1020 nodes, 2615 edges, PR-AUC 0.2348, High band precision 85.19%) ko dimag me lock kar lein.\n"
        "3. Step 3: Section 31 ke 105 Viva Questions aur Section 32 ke 50 Rapid Fire questions ke 'Viva me bolne ka simple answer' ko multiple times practice karein.\n"
        "4. Step 4: Section 38 ka live demonstration script open laptop par click karte huye rehears karein.")

    doc.add_page_break()

def main():
    print("[*] Initializing Master Viva Guide Document Generation...")
    doc = setup_document()

    print("[*] Generating Prestige Cover Page...")
    build_cover_page(doc)

    print("[*] Generating Table of Contents Roadmap...")
    build_table_of_contents_summary(doc)

    print("[*] Assembling Part 1 (Sections 1 to 15)...")
    build_part1(doc)
    doc.add_page_break()

    print("[*] Assembling Part 2 (Sections 16 to 30)...")
    build_part2(doc)
    doc.add_page_break()

    print("[*] Assembling Part 3 (Sections 31 to 35)...")
    build_part3(doc)
    doc.add_page_break()

    print("[*] Assembling Part 4 (Sections 36 to 41)...")
    build_part4(doc)

    out_path = os.path.join(os.path.dirname(SCRIPT_DIR), "PROJECT_VIVA_PREPARATION_GUIDE.docx")
    print(f"[*] Saving master document to: {out_path}")
    doc.save(out_path)
    
    file_size_mb = os.path.getsize(out_path) / (1024 * 1024)
    para_count = len(doc.paragraphs)
    table_count = len(doc.tables)
    print(f"[+] SUCCESS! File generated successfully.")
    print(f"    - File Size: {file_size_mb:.2f} MB")
    print(f"    - Paragraphs: {para_count}")
    print(f"    - Tables: {table_count}")

if __name__ == "__main__":
    main()
